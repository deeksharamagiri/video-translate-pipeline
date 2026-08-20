"""
Generates the always-created Job Report DOCX.

The report contains:

- Job metadata
- Pre-processing information
- Automatic translation quality
- Translation quality components
- Translation warnings
- Voiceover quality
- Numeric/decimal TTS metrics
- Per-segment translation
- Per-segment voiceover QC

Important:

The Automatic Translation Quality score is a proxy.
It does not represent human-verified semantic translation
accuracy because no reference translation is available.
"""

from typing import List, Optional, Dict

from docx import Document
from docx.shared import RGBColor

from pipeline.stage3_segment_tm import Segment


LOW_CONFIDENCE_THRESHOLD = 0.55


def _quality_label(
    score: float,
) -> str:

    if score >= 90:
        return "EXCELLENT"

    if score >= 75:
        return "GOOD"

    if score >= 60:
        return "FAIR"

    return "NEEDS REVIEW"


def _add_score_paragraph(
    doc,
    title: str,
    score: float,
):
    p = doc.add_paragraph()

    run = p.add_run(
        f"{title}: "
        f"{score:.1f}/100 — "
        f"{_quality_label(score)}"
    )

    run.bold = True

    if score < 75:

        run.font.color.rgb = RGBColor(
            0xC0,
            0x50,
            0x00,
        )

    return p


def generate_job_report(
    job_id: str,
    segments: List[Segment],
    source_lang: str,
    target_lang: str,
    engine_name: str,
    preprocess_info: dict,
    out_docx_path: str,
    quality_info: Optional[Dict] = None,
) -> str:

    doc = Document()

    doc.add_heading(
        "Video/Audio Translation — Job Report",
        level=1,
    )

    # =====================================================
    # Metadata
    # =====================================================

    meta = doc.add_paragraph()

    meta.add_run(
        "Job ID: "
    ).bold = True

    meta.add_run(
        f"{job_id}\n"
    )

    meta.add_run(
        "Source language: "
    ).bold = True

    meta.add_run(
        f"{source_lang}\n"
    )

    meta.add_run(
        "Target language: "
    ).bold = True

    meta.add_run(
        f"{target_lang}\n"
    )

    meta.add_run(
        "Translation engine: "
    ).bold = True

    meta.add_run(
        f"{engine_name}\n"
    )

    meta.add_run(
        "Total segments: "
    ).bold = True

    meta.add_run(
        f"{len(segments)}\n"
    )

    low_conf = [
        s
        for s in segments
        if s.confidence
        < LOW_CONFIDENCE_THRESHOLD
    ]

    meta.add_run(
        "Low-confidence segments flagged: "
    ).bold = True

    meta.add_run(
        f"{len(low_conf)}"
    )

    # =====================================================
    # Automatic Quality Assessment
    # =====================================================

    if quality_info:

        doc.add_heading(
            "Automatic Quality Assessment",
            level=2,
        )

        doc.add_paragraph(
            "The report contains separate "
            "Translation Quality and Voiceover "
            "Quality scores."
        )

        doc.add_paragraph(
            "Important: Automatic Translation Quality "
            "is a quality proxy based on ASR confidence, "
            "translation completeness, and numeric "
            "preservation. It is not a human-verified "
            "semantic translation-accuracy score."
        )

        # -------------------------------------------------
        # Translation Quality
        # -------------------------------------------------

        doc.add_heading(
            "Automatic Translation Quality",
            level=3,
        )

        translation_score = (
            quality_info.get(
                "overall_translation_quality_score",
                0,
            )
        )

        _add_score_paragraph(
            doc,
            "Overall Translation Quality",
            translation_score,
        )

        translation_table = doc.add_table(
            rows=1,
            cols=3,
        )

        translation_table.style = (
            "Light Grid Accent 1"
        )

        headers = (
            translation_table
            .rows[0]
            .cells
        )

        headers[0].text = "Metric"
        headers[1].text = "Score"
        headers[2].text = "Interpretation"

        translation_metrics = [
            (
                "ASR confidence",
                quality_info.get(
                    "average_asr_confidence",
                    0,
                ),
                "Average confidence of the speech-recognition transcript.",
            ),
            (
                "Translation completeness",
                quality_info.get(
                    "translation_completeness_score",
                    0,
                ),
                "Checks for missing or suspiciously short/long translations.",
            ),
            (
                "Numeric preservation",
                quality_info.get(
                    "numeric_preservation_score",
                    0,
                ),
                "Checks whether numeric expressions are preserved in translation.",
            ),
        ]

        for (
            name,
            value,
            explanation,
        ) in translation_metrics:

            row = (
                translation_table
                .add_row()
                .cells
            )

            row[0].text = name
            row[1].text = (
                f"{value:.1f}/100"
            )
            row[2].text = explanation

        doc.add_paragraph(
            "Translation Quality weighting: "
            "40% ASR confidence + "
            "35% translation completeness + "
            "25% numeric preservation."
        )

        # -------------------------------------------------
        # Translation warnings
        # -------------------------------------------------

        translation_warnings = (
            quality_info.get(
                "translation_warnings",
                [],
            )
        )

        if translation_warnings:

            doc.add_heading(
                "Translation Quality Warnings",
                level=4,
            )

            for warning in (
                translation_warnings
            ):

                warn_p = (
                    doc.add_paragraph()
                )

                run = warn_p.add_run(
                    f"⚠ {warning}"
                )

                run.font.color.rgb = (
                    RGBColor(
                        0xC0,
                        0x50,
                        0x00,
                    )
                )

        # -------------------------------------------------
        # Voiceover Quality
        # -------------------------------------------------

        doc.add_heading(
            "Automatic Voiceover Quality",
            level=3,
        )

        score = quality_info.get(
            "overall_quality_score",
            0,
        )

        _add_score_paragraph(
            doc,
            "Overall Voiceover Quality",
            score,
        )

        quality_table = doc.add_table(
            rows=1,
            cols=3,
        )

        quality_table.style = (
            "Light Grid Accent 1"
        )

        headers = (
            quality_table
            .rows[0]
            .cells
        )

        headers[0].text = "Metric"
        headers[1].text = "Score"
        headers[2].text = "Interpretation"

        metrics = [
            (
                "TTS coverage",
                quality_info.get(
                    "tts_coverage_percent",
                    0,
                ),
                "Percentage of segments with generated voiceover.",
            ),
            (
                "Numeric TTS normalization",
                quality_info.get(
                    "numeric_normalization_percent",
                    100,
                ),
                "Percentage of detected numeric expressions successfully prepared for speech.",
            ),
            (
                "Timing alignment",
                quality_info.get(
                    "average_alignment_score",
                    0,
                ),
                "How closely generated speech duration matches each segment slot.",
            ),
        ]

        for (
            name,
            value,
            explanation,
        ) in metrics:

            row = (
                quality_table
                .add_row()
                .cells
            )

            row[0].text = name
            row[1].text = (
                f"{value:.1f}/100"
            )
            row[2].text = explanation

        doc.add_paragraph(
            "Voiceover Quality weighting: "
            "35% TTS coverage + "
            "30% numeric normalization + "
            "25% timing alignment + "
            "10% base score, minus TTS failure penalty."
        )

        # -------------------------------------------------
        # Numeric statistics
        # -------------------------------------------------

        doc.add_heading(
            "Numeric / Decimal Speech QC",
            level=3,
        )

        numeric_p = doc.add_paragraph()

        numeric_p.add_run(
            "Segments containing numbers: "
        ).bold = True

        numeric_p.add_run(
            str(
                quality_info.get(
                    "numeric_segments",
                    0,
                )
            )
        )

        numeric_p.add_run(
            "\nNumeric expressions detected: "
        ).bold = True

        numeric_p.add_run(
            str(
                quality_info.get(
                    "numeric_expressions",
                    0,
                )
            )
        )

        numeric_p.add_run(
            "\nDecimal expressions detected: "
        ).bold = True

        numeric_p.add_run(
            str(
                quality_info.get(
                    "decimal_expressions",
                    0,
                )
            )
        )

        numeric_p.add_run(
            "\nNormalization successes: "
        ).bold = True

        numeric_p.add_run(
            str(
                quality_info.get(
                    "numeric_normalization_success",
                    0,
                )
            )
        )

        numeric_p.add_run(
            "\nNormalization failures: "
        ).bold = True

        numeric_p.add_run(
            str(
                quality_info.get(
                    "numeric_normalization_failures",
                    0,
                )
            )
        )

        # -------------------------------------------------
        # TTS failures
        # -------------------------------------------------

        tts_failed = quality_info.get(
            "tts_failed_segments",
            0,
        )

        tts_success = quality_info.get(
            "tts_success_segments",
            0,
        )

        doc.add_heading(
            "Voiceover Generation",
            level=3,
        )

        tts_p = doc.add_paragraph()

        tts_p.add_run(
            "Successful TTS segments: "
        ).bold = True

        tts_p.add_run(
            f"{tts_success}\n"
        )

        tts_p.add_run(
            "Failed TTS segments: "
        ).bold = True

        tts_p.add_run(
            f"{tts_failed}"
        )

        # -------------------------------------------------
        # Voiceover warnings
        # -------------------------------------------------

        warnings = quality_info.get(
            "warnings",
            [],
        )

        if warnings:

            doc.add_heading(
                "Voiceover Quality Warnings",
                level=3,
            )

            for warning in warnings:

                warn_p = (
                    doc.add_paragraph()
                )

                run = warn_p.add_run(
                    f"⚠ {warning}"
                )

                run.font.color.rgb = (
                    RGBColor(
                        0xC0,
                        0x50,
                        0x00,
                    )
                )

    # =====================================================
    # Pre-processing
    # =====================================================

    doc.add_heading(
        "Pre-Processing Notes",
        level=2,
    )

    p = doc.add_paragraph()

    p.add_run(
        f"Input type: "
        f"{preprocess_info.get('input_kind')}\n"
    )

    p.add_run(
        f"Subtitle track found: "
        f"{preprocess_info.get('subtitle_track_found')}\n"
    )

    p.add_run(
        f"Denoise applied: "
        f"{preprocess_info.get('denoise_applied')}\n"
    )

    if preprocess_info.get(
        "snr_db"
    ) is not None:

        p.add_run(
            f"Estimated SNR: "
            f"{preprocess_info.get('snr_db')} dB\n"
        )

    for warning in preprocess_info.get(
        "warnings",
        [],
    ):

        warn_p = doc.add_paragraph()

        run = warn_p.add_run(
            f"⚠ {warning}"
        )

        run.font.color.rgb = RGBColor(
            0xC0,
            0x50,
            0x00,
        )

    # =====================================================
    # Transcript + translation
    # =====================================================

    doc.add_heading(
        "Full Transcript & Translation",
        level=2,
    )

    table = doc.add_table(
        rows=1,
        cols=5,
    )

    table.style = (
        "Light Grid Accent 1"
    )

    hdr = table.rows[0].cells

    hdr[0].text = "#"
    hdr[1].text = "Time"
    hdr[2].text = "Source Text"
    hdr[3].text = "Translated Text"
    hdr[4].text = "Confidence / Flag"

    for i, seg in enumerate(
        segments,
        start=1,
    ):

        row = table.add_row().cells

        row[0].text = str(i)

        row[1].text = (
            f"{_fmt_time(seg.start)}"
            f" - "
            f"{_fmt_time(seg.end)}"
        )

        row[2].text = (
            seg.text or ""
        )

        row[3].text = (
            seg.translated_text
            or ""
        )

        flag = (
            "⚠ LOW CONFIDENCE"
            if seg.confidence
            < LOW_CONFIDENCE_THRESHOLD
            else "OK"
        )

        cache_note = (
            " (cache)"
            if seg.from_cache
            else ""
        )

        row[4].text = (
            f"{seg.confidence:.2f}"
            f" — {flag}"
            f"{cache_note}"
        )

    # =====================================================
    # Per-segment voiceover QC
    # =====================================================

    if quality_info:

        details = quality_info.get(
            "segment_details",
            [],
        )

        if details:

            doc.add_heading(
                "Voiceover Segment QC",
                level=2,
            )

            table = doc.add_table(
                rows=1,
                cols=6,
            )

            table.style = (
                "Light Grid Accent 1"
            )

            hdr = table.rows[0].cells

            hdr[0].text = "Segment"
            hdr[1].text = "Status"
            hdr[2].text = "Slot (sec)"
            hdr[3].text = "TTS (sec)"
            hdr[4].text = "Alignment"
            hdr[5].text = "Numbers"

            for item in details:

                row = (
                    table
                    .add_row()
                    .cells
                )

                row[0].text = str(
                    item.get(
                        "segment",
                        "",
                    )
                )

                row[1].text = str(
                    item.get(
                        "status",
                        "",
                    )
                )

                row[2].text = (
                    f"{item.get('slot_duration', 0):.2f}"
                )

                row[3].text = (
                    f"{item.get('tts_duration', 0):.2f}"
                )

                row[4].text = (
                    f"{item.get('alignment_score', 0):.1f}"
                )

                row[5].text = str(
                    item.get(
                        "numeric_count",
                        0,
                    )
                )

    # =====================================================
    # Save
    # =====================================================

    doc.save(
        out_docx_path
    )

    return out_docx_path


def _fmt_time(
    seconds: float,
) -> str:

    m, s = divmod(
        int(seconds),
        60,
    )

    h, m = divmod(
        m,
        60,
    )

    return (
        f"{h:02d}:"
        f"{m:02d}:"
        f"{s:02d}"
    )